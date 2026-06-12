"""Export tests for NBCOT."""

from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from nbcot.exports import build_orders_workbook
from nbcot.tests import fixtures


class CiscoOrderExportTest(TestCase):
    """Test tracked order Excel exports."""

    def test_export_contains_order_fields_and_major_minor_line_classification(self):
        """Workbook should expose order metadata and classify line hierarchy."""
        order = fixtures.create_ciscoorder(
            order_number="SO-9100",
            project_number="PRJ-9100",
            notes="Important rollout",
            notification_recipients="ops@example.com",
        )
        fixtures.create_line(order, line_key="major", line_number="1.0", sku="MAJOR-SKU", is_tracked=True)
        fixtures.create_line(
            order,
            line_key="minor",
            line_number="1.0.1",
            sku="MINOR-SKU",
            serial_number="SER-9100",
            tracking_number="TRACK-9100",
            tracking_url="https://carrier.example/TRACK-9100",
        )

        content = build_orders_workbook([order])
        workbook = load_workbook(BytesIO(content))

        self.assertEqual(workbook.sheetnames, ["Tracked Orders", "SO-9100"])
        order_sheet = workbook["Tracked Orders"]
        line_sheet = workbook["SO-9100"]
        self.assertEqual(order_sheet["A2"].value, "SO-9100")
        self.assertEqual(order_sheet["A2"].hyperlink.location, "'SO-9100'!A1")
        self.assertEqual(order_sheet["D2"].value, "PRJ-9100")
        self.assertEqual(order_sheet["E2"].value, "Important rollout")
        self.assertEqual(line_sheet["E2"].value, "Major")
        self.assertEqual(line_sheet["E3"].value, "Minor")
        self.assertEqual(line_sheet["J3"].value, "SER-9100")
        self.assertEqual(line_sheet["N3"].value, "TRACK-9100")

    def test_export_creates_one_sheet_per_order_with_unique_safe_names(self):
        """Multi-order workbooks should use a linked summary and one sheet per order."""
        first_order = fixtures.create_ciscoorder(order_number="SO/9200")
        second_order = fixtures.create_ciscoorder(order_number="SO/9200", environment="prod")
        fixtures.create_line(first_order, line_key="major", line_number="1.0", sku="FIRST")
        fixtures.create_line(second_order, line_key="major", line_number="1.0", sku="SECOND")

        workbook = load_workbook(BytesIO(build_orders_workbook([first_order, second_order])))

        self.assertEqual(workbook.sheetnames, ["Tracked Orders", "SO-9200", "SO-9200-2"])
        self.assertEqual(workbook["Tracked Orders"]["A2"].hyperlink.location, "'SO-9200'!A1")
        self.assertEqual(workbook["Tracked Orders"]["A3"].hyperlink.location, "'SO-9200-2'!A1")
        self.assertEqual(workbook["SO-9200"]["F2"].value, "FIRST")
        self.assertEqual(workbook["SO-9200-2"]["F2"].value, "SECOND")
