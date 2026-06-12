"""Serial attribute helper tests for NBCOT."""

from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from nbcot.cisco.serials import (
    joined_attribute_values,
    serial_attribute_values,
    serial_values_display,
)
from nbcot.exports import build_order_workbook


class LineList(list):
    """Small stand-in for a Django related manager."""

    def all(self):
        """Return all fake line objects."""
        return self


def _line(**kwargs):
    defaults = {
        "line_key": "phones",
        "line_number": "3.0",
        "line_sort_key": "0003.0000",
        "sku": "CP-7841-K9=",
        "description": "Cisco IP Phone 7841",
        "status": "Closed",
        "shipment_status": "Closed",
        "quantity_ordered": 30,
        "quantity_fulfilled": 30,
        "quantity_backordered": 0,
        "promised_delivery_date": None,
        "estimated_delivery_date": None,
        "actual_delivery_date": None,
        "estimated_ship_date": None,
        "serial_number": "",
        "parent_serial_number": "",
        "mac_address": "",
        "imei_number": "",
        "instance_number": "",
        "license_key": "",
        "cloud_id": "",
        "contract_number": "",
        "ship_set": "1",
        "carrier": "DHL",
        "tracking_number": "TRACK-1",
        "tracking_url": "",
        "proof_of_delivery_url": "",
        "is_tracked": False,
        "raw_payload": {},
    }
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def _order(lines):
    return SimpleNamespace(
        order_number="SO-9050",
        environment="prod",
        customer_po_number="PO-9050",
        project_number="",
        notes="",
        account_name="Acme Corp",
        account_number="",
        status="ORDER_IN_PROGRESS",
        status_detail="",
        lifecycle_state="",
        is_tracked=True,
        is_archived=False,
        requested_delivery_date=None,
        promised_delivery_date=None,
        estimated_delivery_date=None,
        ordered_at=None,
        last_event_at=None,
        last_synced_at=None,
        last_sync_status="Success",
        notification_recipients="",
        web_order_url="",
        cisco_sales_order_url="",
        lines=LineList(lines),
    )


def test_serial_attribute_values_include_all_serials_and_macs():
    """Helpers should expose every Cisco serial attribute, not only the first value."""
    raw_payload = {
        "serialNumberAttributes": [
            {
                "serialNumber": ["SER-001", "SER-002"],
                "macAddresses": ["MAC-001", "MAC-002"],
            },
            {
                "serialNumber": "SER-003",
                "macAddresses": "MAC-003",
            },
        ]
    }

    assert serial_attribute_values(raw_payload, "serialNumber") == ["SER-001", "SER-002", "SER-003"]
    assert serial_attribute_values(raw_payload, "macAddresses") == ["MAC-001", "MAC-002", "MAC-003"]
    assert joined_attribute_values(raw_payload, "serialNumber") == "SER-001\nSER-002\nSER-003"


def test_serial_values_display_caps_large_ui_lists_but_export_keeps_all_values():
    """UI display should be bounded while export helpers retain all values."""
    serials = [f"SER-{index:04d}" for index in range(1, 31)]
    display = serial_values_display(serials)

    assert "SER-0001" in display
    assert "SER-0010" in display
    assert "SER-0011" not in display
    assert "+20 more" in display
    assert "\n".join(serials).splitlines()[-1] == "SER-0030"


def test_single_order_export_uses_cisco_like_sheet_names_and_full_serial_lists():
    """Single-order workbook should mimic Cisco's line detail shape and include all serial values."""
    line = _line(
        serial_number="SER-001",
        mac_address="MAC-001",
        raw_payload={
            "serialNumberAttributes": [
                {
                    "serialNumber": ["SER-001", "SER-002"],
                    "macAddresses": ["MAC-001", "MAC-002"],
                },
                {
                    "serialNumber": ["SER-003"],
                    "macAddresses": ["MAC-003"],
                },
            ]
        },
    )

    workbook = load_workbook(BytesIO(build_order_workbook(_order([line]))))

    assert workbook.sheetnames == ["Order Header", "Order Line Details", "Credit Breakdown"]
    line_sheet = workbook["Order Line Details"]
    assert line_sheet["A1"].value == "Line Number"
    assert line_sheet["M1"].value == "Serial Numbers"
    assert line_sheet["A2"].value == "3.0"
    assert line_sheet["B2"].value == "CP-7841-K9="
    assert line_sheet["M2"].value == "SER-001\nSER-002\nSER-003"
    assert line_sheet.cell(row=1, column=171).value == "MAC Addresses"
    assert line_sheet.cell(row=2, column=171).value == "MAC-001\nMAC-002\nMAC-003"
