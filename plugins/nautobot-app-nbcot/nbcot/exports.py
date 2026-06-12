"""Excel exports for NBCOT tracked orders."""

from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from nbcot.cisco.line_items import build_line_tree


ORDER_HEADERS = [
    "Order Number",
    "Environment",
    "Customer PO Number",
    "Project Number",
    "Notes",
    "Account Name",
    "Account Number",
    "Status",
    "Status Detail",
    "Lifecycle State",
    "Is Tracked",
    "Is Archived",
    "Requested Delivery Date",
    "Promised Delivery Date",
    "Estimated Delivery Date",
    "Last Synced At",
    "Last Sync Status",
    "Notification Recipients",
]

LINE_HEADERS = [
    "Order Number",
    "Line Number",
    "Parent Line",
    "Line Key",
    "Classification",
    "SKU",
    "Description",
    "Status",
    "Shipment Status",
    "Serial Number",
    "MAC Address",
    "Instance Number",
    "Carrier",
    "Tracking Number",
    "Tracking URL",
    "Proof Of Delivery URL",
    "Quantity Ordered",
    "Quantity Fulfilled",
    "Quantity Backordered",
    "Promised Delivery Date",
    "Estimated Delivery Date",
    "Actual Delivery Date",
    "Estimated Ship Date",
    "Is Tracked",
]


def _date_value(value):
    return value.isoformat() if value else ""


def _style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F81BD")


def _autosize(sheet):
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width, 12), 60)


def _classification(row) -> str:
    return "Major" if row.depth == 0 else "Minor"


def build_orders_workbook(orders: Iterable) -> bytes:
    """Build an XLSX workbook for the supplied Cisco orders."""
    workbook = Workbook()
    orders_sheet = workbook.active
    orders_sheet.title = "Tracked Orders"
    lines_sheet = workbook.create_sheet("Line Items")
    orders_sheet.append(ORDER_HEADERS)
    lines_sheet.append(LINE_HEADERS)
    _style_header(orders_sheet)
    _style_header(lines_sheet)

    for order in orders:
        orders_sheet.append(
            [
                order.order_number,
                order.environment.upper(),
                order.customer_po_number,
                order.project_number,
                order.notes,
                order.account_name,
                order.account_number,
                order.status,
                order.status_detail,
                order.lifecycle_state,
                "Yes" if order.is_tracked else "No",
                "Yes" if order.is_archived else "No",
                _date_value(order.requested_delivery_date),
                _date_value(order.promised_delivery_date),
                _date_value(order.estimated_delivery_date),
                order.last_synced_at.isoformat() if order.last_synced_at else "",
                order.last_sync_status,
                order.notification_recipients,
            ]
        )

        line_rows = build_line_tree(order.lines.all())
        line_by_row_id = {row.row_id: row.line for row in line_rows}
        for row in line_rows:
            line = row.line
            parent_line = ""
            if row.parent_id:
                parent = line_by_row_id.get(row.parent_id)
                if parent is not None:
                    parent_line = parent.line_number
            lines_sheet.append(
                [
                    order.order_number,
                    line.line_number,
                    parent_line,
                    line.line_key,
                    _classification(row),
                    line.sku,
                    line.description,
                    line.status,
                    line.shipment_status,
                    line.serial_number,
                    line.mac_address,
                    line.instance_number,
                    line.carrier,
                    line.tracking_number,
                    line.tracking_url,
                    line.proof_of_delivery_url,
                    line.quantity_ordered,
                    line.quantity_fulfilled,
                    line.quantity_backordered,
                    _date_value(line.promised_delivery_date),
                    _date_value(line.estimated_delivery_date),
                    _date_value(line.actual_delivery_date),
                    _date_value(line.estimated_ship_date),
                    "Yes" if line.is_tracked else "No",
                ]
            )
            if line.is_tracked:
                for cell in lines_sheet[lines_sheet.max_row]:
                    cell.fill = PatternFill("solid", fgColor="EAF4FF")

    _autosize(orders_sheet)
    _autosize(lines_sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
